"""Lectura de documentos fuente y conversión a texto plano para la IA.

Soporta: PDF digital (texto + tablas), correo .msg de Outlook,
Excel, Word e imágenes (estas últimas se envían como visión).
"""
import base64
from pathlib import Path

TEXTO_MAX_CHARS = 60_000


class DocumentoIlegibleError(Exception):
    """El documento no contiene texto extraíble (ej. PDF escaneado sin OCR)."""


def leer_documento(ruta: str) -> dict:
    """Devuelve {"texto": str} o {"imagen_b64": str, "mime": str} según el tipo."""
    path = Path(ruta)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return {"texto": _leer_pdf(path)}
    if ext == ".msg":
        return {"texto": _leer_msg(path)}
    if ext in (".xlsx", ".xls"):
        return {"texto": _leer_excel(path)}
    if ext == ".docx":
        return {"texto": _leer_word(path)}
    if ext in (".jpg", ".jpeg", ".png"):
        return _leer_imagen(path)
    raise DocumentoIlegibleError(f"Extensión no soportada: {ext}")


ANCHO_MINIMO_IMAGEN = 1400


def _leer_imagen(path: Path) -> dict:
    """Prepara la imagen para visión. Las imágenes pequeñas se reescalan:
    con poca resolución el modelo confunde dígitos en celdas de tablas."""
    import io

    from PIL import Image

    img = Image.open(path)
    if img.width < ANCHO_MINIMO_IMAGEN:
        factor = ANCHO_MINIMO_IMAGEN / img.width
        img = img.convert("RGB").resize(
            (ANCHO_MINIMO_IMAGEN, round(img.height * factor)), Image.LANCZOS
        )
    buffer = io.BytesIO()
    img.convert("RGB").save(buffer, format="PNG")
    b64 = base64.b64encode(buffer.getvalue()).decode()
    return {"imagen_b64": b64, "mime": "image/png"}


def _leer_pdf(path: Path) -> str:
    import pdfplumber

    partes: list[str] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            texto = page.extract_text() or ""
            partes.append(f"--- PÁGINA {i + 1} ---\n{texto}")
            # Las tablas se agregan aparte: pdfplumber preserva la estructura
            # de celdas mejor que el texto lineal.
            for ti, tabla in enumerate(page.extract_tables()):
                filas = [
                    " | ".join("" if c is None else str(c).replace("\n", " ") for c in fila)
                    for fila in tabla
                ]
                partes.append(f"[TABLA {ti + 1} PÁGINA {i + 1}]\n" + "\n".join(filas))

    contenido = "\n\n".join(partes)
    if len(contenido.replace("-", "").replace("PÁGINA", "").strip()) < 50:
        raise DocumentoIlegibleError(
            "El PDF no contiene texto extraíble; parece un documento escaneado."
        )
    return contenido[:TEXTO_MAX_CHARS]


def _leer_msg(path: Path) -> str:
    import extract_msg

    msg = extract_msg.Message(str(path))
    cuerpo = msg.body or ""
    encabezado = (
        f"ASUNTO: {msg.subject}\n"
        f"DE: {msg.sender}\n"
        f"FECHA: {msg.date}\n"
        "--- CUERPO DEL CORREO (el mensaje más reciente aparece primero; "
        "las líneas con '>' son correos anteriores citados) ---\n"
    )
    return (encabezado + cuerpo)[:TEXTO_MAX_CHARS]


def _leer_excel(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    partes: list[str] = []
    for ws in wb.worksheets:
        partes.append(f"--- HOJA: {ws.title} ---")
        for fila in ws.iter_rows(values_only=True):
            celdas = ["" if c is None else str(c) for c in fila]
            if any(c.strip() for c in celdas):
                partes.append(" | ".join(celdas))
    wb.close()
    return "\n".join(partes)[:TEXTO_MAX_CHARS]


def _leer_word(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(celda.text.strip() for celda in fila.cells))
    return "\n".join(partes)[:TEXTO_MAX_CHARS]
